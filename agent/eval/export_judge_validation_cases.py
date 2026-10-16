"""
Builds the labelling fixture for run_judge_validation_eval.py from the CBA end-to-end
experiment (fixtures/CBA_News_Retrieval_Evaluation_Matching.xlsx, "Answer Quality (All)"
tab): one case per question - a real production answer plus the evidence it was generated
from.

Per question it picks the answer the experiment's judge scored LEAST grounded: those carry
the most unsupported claims, so they give the most chances to catch the judge both wrongly
rejecting and wrongly accepting a claim.

What it fills in vs. leaves for a human (same rule as export_labels.py: no fabricated labels):
  claims[].claim                 drafted by ANALYSIS_MODEL splitting the answer into atomic
                                 statements - edit/split/merge freely
  claims[].covered_in_evidence   null <- human: does the evidence state it?
  facts[].fact                   drafted by ANALYSIS_MODEL: the question-relevant facts in
                                 THIS case's evidence (what a complete answer should say)
  facts[].relevant               null <- human: is it at least partly relevant to the question?
                                 (also add a few clearly off-topic facts from the evidence, as
                                 relevant=false, to test that the judge doesn't penalise them)
  facts[].covered_in_answer      null <- human: does the answer convey it?

Refuses to overwrite an existing fixture without --force, since that file holds the
human's labels once work starts.

Usage:
    python -m agent.eval.export_judge_validation_cases
    python -m agent.eval.export_judge_validation_cases --output path.json --force
"""

import argparse
import json
import os

import openpyxl
from langchain_anthropic import ChatAnthropic
from pydantic import BaseModel, Field

from agent.shared.news_text import format_news_content
from config import ANALYSIS_MODEL, ANTHROPIC_API_KEY
from db.client import get_client

WORKBOOK = "agent/eval/fixtures/CBA_News_Retrieval_Evaluation_Matching.xlsx"
DEFAULT_OUTPUT = "agent/eval/fixtures/judge_validation_test_set.json"


class Claims(BaseModel):
    claims: list[str] = Field(description="Every atomic factual or interpretive statement the answer makes, "
                              "one per item, close to the answer's own wording. Skip pure formatting/headings.")


class Facts(BaseModel):
    facts: list[str] = Field(description="The distinct facts in the evidence that are RELEVANT to the question - "
                             "what a complete answer should mention. One short sentence each. Empty list if no "
                             "evidence is relevant to the question.")


def _read_experiment(path: str) -> tuple[dict, list[dict]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    row_to_id = {r[0]: r[1] for r in wb["Original Articles"].iter_rows(min_row=2, values_only=True) if r[0]}
    rows = list(wb["Answer Quality (All)"].iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if r[0] == "k" and r[1] == "Question ID")
    cols = {name: i for i, name in enumerate(rows[h])}
    parse = lambda s: [] if s in (None, "(none)") else [int(x) for x in str(s).split(", ")]  # noqa: E731
    answers = [{
        "question_id": r[cols["Question ID"]], "question": r[cols["Question"]],
        "retrieved": parse(r[cols["Retrieved (CSV rows)"]]), "groundedness_pct": r[cols["Groundedness %"]],
        "answer": r[cols["Answer"]],
    } for r in rows[h + 1:] if r[0] is not None]
    return row_to_id, answers


def _evidence_text(row_to_id: dict) -> dict:
    """CSV row -> the exact evidence string search_internal_news_tool produces for that article."""
    data = (get_client().table("ticker_news").select("id, title, snippet, publisher")
            .in_("id", list(row_to_id.values())).execute().data)
    by_id = {d["id"]: d for d in data}
    return {row: f"{format_news_content(by_id[aid]['title'], by_id[aid].get('snippet'))} "
                 f"({by_id[aid].get('publisher') or 'Unknown'})" for row, aid in row_to_id.items()}


def _least_grounded_per_question(answers: list[dict]) -> list[dict]:
    picked = []
    for qid in dict.fromkeys(a["question_id"] for a in answers):
        judged = [a for a in answers if a["question_id"] == qid and isinstance(a["groundedness_pct"], (int, float))]
        picked.append(min(judged, key=lambda a: a["groundedness_pct"]))
    return picked


def build(workbook: str) -> list[dict]:
    row_to_id, answers = _read_experiment(workbook)
    text = _evidence_text(row_to_id)
    llm = ChatAnthropic(model=ANALYSIS_MODEL, api_key=ANTHROPIC_API_KEY)
    claim_llm, fact_llm = llm.with_structured_output(Claims), llm.with_structured_output(Facts)

    cases = []
    for a in _least_grounded_per_question(answers):
        evidence = [{"source": "search_internal_news_tool", "content": text[r]} for r in a["retrieved"]]
        claims = claim_llm.invoke(f"Split this answer into atomic statements:\n\n{a['answer']}").claims
        facts = fact_llm.invoke(
            f"Question: {a['question']}\n\nEvidence:\n" + "\n\n".join(e["content"] for e in evidence)
        ).facts
        cases.append({
            "id": a["question_id"].lower(),
            "question": a["question"],
            "answer": a["answer"],
            "evidence": evidence,
            "claims": [{"claim": c, "covered_in_evidence": None} for c in claims],
            "facts": [{"fact": f, "relevant": None, "covered_in_answer": None} for f in facts],
        })
        print(f"  {cases[-1]['id']}: {len(claims)} claims, {len(facts)} facts")
    return cases


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", default=WORKBOOK)
    parser.add_argument("--output", default=DEFAULT_OUTPUT)
    parser.add_argument("--force", action="store_true", help="Overwrite an existing fixture (destroys its labels).")
    args = parser.parse_args()

    if os.path.exists(args.output) and not args.force:
        raise SystemExit(f"{args.output} already exists - it may hold human labels. Pass --force to overwrite.")
    cases = build(args.workbook)
    with open(args.output, "w", encoding="utf-8") as f:
        json.dump(cases, f, indent=2, ensure_ascii=False)
    print(f"Wrote {len(cases)} cases to {args.output} - fill in every null covered_in_evidence / "
          "covered_in_answer, then run: python -m agent.eval.runners.run_judge_validation_eval")
